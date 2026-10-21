"""Extract readable text from web pages and uploaded documents."""

from __future__ import annotations

import io
import re

from ..config import settings
from ..result import Result

_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/130.0 Safari/537.36"
)
_NOISE_TAGS = ["script", "style", "nav", "footer", "header", "aside", "form", "noscript", "iframe", "svg"]


def _clean(text: str, limit: int | None = None) -> tuple[str, bool]:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    limit = limit or settings.max_input_chars
    return text[:limit], len(text) > limit


def html_to_text(html: str) -> tuple[str, str]:
    """Return (title, main text) from raw HTML."""
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else "Untitled page"
    for tag in soup(_NOISE_TAGS):
        tag.decompose()
    root = soup.find("article") or soup.find("main") or soup.body or soup
    blocks = [el.get_text(" ", strip=True) for el in root.find_all(["h1", "h2", "h3", "h4", "p", "li"])]
    text = "\n".join(b for b in blocks if len(b) > 1)
    return title, text or root.get_text(" ", strip=True)


def fetch_url(url: str) -> Result:
    import requests

    if not re.match(r"^https?://", url.strip()):
        return Result.fail("Invalid URL — it must start with http:// or https://")
    try:
        resp = requests.get(url.strip(), headers={"User-Agent": _UA}, timeout=settings.request_timeout_s)
        resp.raise_for_status()
    except requests.exceptions.Timeout:
        return Result.fail("Request timed out — the website took too long to respond.")
    except requests.exceptions.HTTPError as e:
        return Result.fail(f"HTTP error {e.response.status_code}.")
    except requests.exceptions.RequestException as e:
        return Result.fail(f"Could not fetch the page: {e}")

    title, text = html_to_text(resp.text)
    text, truncated = _clean(text)
    if not text:
        return Result.fail("No readable text found on the page.")
    return Result.success(text, title=title, truncated=truncated)


def extract_pdf(data: bytes) -> Result:
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = pdf.pages[: settings.max_pdf_pages]
            text = "\n\n".join((p.extract_text() or "").strip() for p in pages)
            total = len(pdf.pages)
    except Exception as e:  # noqa: BLE001
        return Result.fail(f"PDF extraction error: {e}")
    text, truncated = _clean(text)
    if not text:
        return Result.fail("No extractable text — the PDF may be scanned images.")
    return Result.success(text, pages=total, truncated=truncated or total > settings.max_pdf_pages)


def extract_docx(data: bytes) -> Result:
    try:
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
    except Exception as e:  # noqa: BLE001
        return Result.fail(f"DOCX extraction error: {e}")
    text, truncated = _clean("\n".join(parts))
    return Result.success(text, truncated=truncated) if text else Result.fail("No text found in the document.")


def extract_txt(data: bytes) -> Result:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            raw = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    text, truncated = _clean(raw)
    return Result.success(text, truncated=truncated) if text else Result.fail("The text file is empty.")


def extract_file(name: str, data: bytes) -> Result:
    ext = name.rsplit(".", 1)[-1].lower()
    handler = {"pdf": extract_pdf, "docx": extract_docx, "txt": extract_txt, "md": extract_txt}.get(ext)
    if handler is None:
        return Result.fail(f"Unsupported file type: .{ext}")
    return handler(data)
